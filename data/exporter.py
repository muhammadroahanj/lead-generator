"""Buffered CSV export plus an XLSX roll-up of the final results."""

import csv
import logging
from datetime import datetime
from pathlib import Path

from config import settings
from config.paths import OutputPaths, default_paths
from data.models import Business

logger = logging.getLogger(__name__)

CSV_HEADERS = [
    "Name", "Category", "Address", "Phone", "Business Email",
    "Rating", "Reviews", "Has Website", "Website URL",
    "Facebook", "Instagram", "LinkedIn",
    "Claimed", "Hours", "Price Level",
    "Latitude", "Longitude", "Plus Code", "Place ID", "Google Maps URL",
    "City", "Area", "Metro Area", "State",
    "Search Query", "Scraped At",
]


def _read_header(filepath: Path) -> list[str] | None:
    """Return the header row of an existing CSV, or None if there isn't one."""
    if not filepath.exists() or filepath.stat().st_size == 0:
        return None
    try:
        with open(filepath, "r", newline="", encoding="utf-8") as handle:
            return next(csv.reader(handle), None)
    except (OSError, StopIteration):
        return None


def _rotate_if_stale(filepath: Path) -> None:
    """Move an existing CSV aside when its header no longer matches.

    Appending new-schema rows onto an old-schema file silently misaligns every
    column, which is worse than starting a new file.
    """
    header = _read_header(filepath)
    if header is None or header == CSV_HEADERS:
        return
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    archived = filepath.with_name(f"{filepath.stem}.legacy-{stamp}{filepath.suffix}")
    filepath.rename(archived)
    logger.warning(
        f"{filepath.name} used an older column set; archived it as {archived.name} "
        "and started a new file."
    )


def _ensure_csv(filepath: Path) -> None:
    filepath.parent.mkdir(parents=True, exist_ok=True)
    _rotate_if_stale(filepath)
    if not filepath.exists() or filepath.stat().st_size == 0:
        with open(filepath, "w", newline="", encoding="utf-8") as handle:
            csv.DictWriter(handle, fieldnames=CSV_HEADERS).writeheader()


class CsvWriter:
    """Appends businesses to a CSV, flushing in batches.

    The previous implementation reopened the file for every single record,
    blocking the event loop on synchronous I/O ~120 times per search.
    """

    def __init__(self, filepath: Path, flush_every: int | None = None):
        self.filepath = filepath
        self.flush_every = flush_every if flush_every is not None else settings.CSV_FLUSH_EVERY
        self._buffer: list[dict] = []
        self._written = 0

    def add(self, businesses: list[Business] | Business) -> None:
        if isinstance(businesses, Business):
            businesses = [businesses]
        self._buffer.extend(b.to_csv_dict() for b in businesses)
        if len(self._buffer) >= max(1, self.flush_every):
            self.flush()

    def flush(self) -> None:
        if not self._buffer:
            return
        _ensure_csv(self.filepath)
        with open(self.filepath, "a", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=CSV_HEADERS)
            writer.writerows(self._buffer)
        self._written += len(self._buffer)
        logger.debug(f"Flushed {len(self._buffer)} rows to {self.filepath.name}")
        self._buffer.clear()

    @property
    def pending(self) -> int:
        return len(self._buffer)


class Exporter:
    """Owns the raw + qualified CSV writers for a run."""

    def __init__(self, paths: OutputPaths | None = None, flush_every: int | None = None):
        self.paths = paths or default_paths()
        self.paths.ensure()
        self.raw = CsvWriter(self.paths.leads_raw_csv, flush_every)
        self.qualified = CsvWriter(self.paths.leads_qualified_csv, flush_every)

    def add_raw(self, businesses) -> None:
        self.raw.add(businesses)

    def add_qualified(self, businesses) -> None:
        self.qualified.add(businesses)

    def flush(self) -> None:
        self.raw.flush()
        self.qualified.flush()

    def export_xlsx(self) -> Path | None:
        """Write both CSVs into a two-sheet workbook. Best-effort."""
        try:
            from openpyxl import Workbook
        except ImportError:
            logger.debug("openpyxl not installed; skipping XLSX export.")
            return None

        self.flush()
        sources = [
            ("Qualified Leads", self.paths.leads_qualified_csv),
            ("All Businesses", self.paths.leads_raw_csv),
        ]
        if not any(path.exists() for _, path in sources):
            return None

        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name, path in sources:
            if not path.exists():
                continue
            sheet = workbook.create_sheet(title=sheet_name)
            with open(path, "r", newline="", encoding="utf-8") as handle:
                for row in csv.reader(handle):
                    sheet.append(row)
            sheet.freeze_panes = "A2"

        if not workbook.sheetnames:
            return None
        try:
            workbook.save(self.paths.leads_xlsx)
        except OSError as exc:
            # Commonly: the workbook is open in Excel and therefore locked.
            logger.warning(f"Could not write {self.paths.leads_xlsx.name}: {exc}")
            return None
        return self.paths.leads_xlsx


def get_csv_row_count(filepath: Path) -> int:
    """Count data rows in a CSV (excluding the header)."""
    if not filepath.exists():
        return 0
    with open(filepath, "r", newline="", encoding="utf-8") as handle:
        return max(0, sum(1 for _ in csv.reader(handle)) - 1)
